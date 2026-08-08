#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Difference de contenu entre deux passes d'un meme document.

Le controle qu'une boucle de correction automatisee doit executer avant de
livrer, et dont la sortie humaine se met directement dans le corps du mail de
livraison.

Methode : ouvrir reellement les deux documents, en extraire les blocs de texte
dans l'ordre de lecture (formes et cellules de tableau comprises), normaliser
(accents, casse, espaces), comparer par appariement de sequences.

C'est ce controle qui a montre qu'une version annoncee « corrigee » ne differait
de la precedente que par deux blocs, aucun ne concernant la demande du client, et
que ces deux blocs reparaient en silence un defaut de la version anterieure.

Trois usages, dans l'ordre d'importance :
  1. refuser de livrer si la zone visee par la demande n'a pas bouge ;
  2. lister dans le mail de livraison ce qui a effectivement change ;
  3. detecter le cas « demande deja satisfaite » (diff vide sur la zone visee).

Usage :
    block_diff.py <avant.pptx> <apres.pptx> [--json <sortie.json>]
Sortie humaine sur stdout, JSON structure si --json.
Code retour : 0 s'il y a des differences, 3 si les deux fichiers sont identiques.

Licence : MIT (voir LICENSE-CODE).
"""
import difflib
import json
import os
import re
import sys
import unicodedata


def sansaccent(s):
    return "".join(c for c in unicodedata.normalize("NFD", s or "")
                   if unicodedata.category(c) != "Mn").lower()


def normaliser(s):
    return re.sub(r"\s+", " ", sansaccent(s)).strip()


def blocs(chemin):
    """[(page, texte)] dans l'ordre de lecture du document."""
    from pptx import Presentation
    prs = Presentation(chemin)
    out = []
    for i, slide in enumerate(prs.slides, start=1):
        for sh in slide.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                out.append((i, sh.text_frame.text.strip()))
            if getattr(sh, "has_table", False) and sh.has_table:
                for row in sh.table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            out.append((i, cell.text.strip()))
    return out


def comparer(avant, apres):
    ba, bb = blocs(avant), blocs(apres)
    na = [normaliser(t) for _, t in ba]
    nb = [normaliser(t) for _, t in bb]
    sm = difflib.SequenceMatcher(a=na, b=nb, autojunk=False)
    res = {"avant": os.path.basename(avant), "apres": os.path.basename(apres),
           "blocs_avant": len(ba), "blocs_apres": len(bb),
           "modifies": [], "ajoutes": [], "retires": []}
    for op, i1, i2, j1, j2 in sm.get_opcodes():
        if op == "equal":
            continue
        if op == "replace":
            # Appariement dans l'ordre : c'est le cas courant d'un bloc reecrit.
            for k in range(max(i2 - i1, j2 - j1)):
                av = ba[i1 + k] if i1 + k < i2 else None
                ap = bb[j1 + k] if j1 + k < j2 else None
                if av and ap:
                    res["modifies"].append({"page_avant": av[0], "page_apres": ap[0],
                                            "avant": av[1], "apres": ap[1]})
                elif av:
                    res["retires"].append({"page": av[0], "texte": av[1]})
                elif ap:
                    res["ajoutes"].append({"page": ap[0], "texte": ap[1]})
        elif op == "delete":
            for k in range(i1, i2):
                res["retires"].append({"page": ba[k][0], "texte": ba[k][1]})
        elif op == "insert":
            for k in range(j1, j2):
                res["ajoutes"].append({"page": bb[k][0], "texte": bb[k][1]})
    res["nb_differences"] = len(res["modifies"]) + len(res["ajoutes"]) + len(res["retires"])
    return res


def resume(t, n=160):
    t = re.sub(r"\s+", " ", t).strip()
    return t if len(t) <= n else t[:n - 1] + "…"


def rendu_humain(res):
    """Texte destine au mail du client : ce qui a change, en clair."""
    lignes = []
    for m in res["modifies"]:
        lignes.append("page %s : « %s » devient « %s »"
                      % (m["page_apres"], resume(m["avant"]), resume(m["apres"])))
    for a in res["ajoutes"]:
        lignes.append("page %s, ajout : « %s »" % (a["page"], resume(a["texte"])))
    for r in res["retires"]:
        lignes.append("page %s, retrait : « %s »" % (r["page"], resume(r["texte"])))
    return lignes


def zone_a_bouge(res, pages=None, motif=None, exclure=None):
    """La zone visee par la demande du client a-t-elle change ?

    Renvoie False si aucune difference retenue ne tombe dans la zone, ce qui doit
    empecher la livraison telle quelle.

    pages   : numeros de page a considerer. None = toutes.
    motif   : expression reguliere que le texte de la difference doit contenir.
    exclure : expression reguliere qui DISQUALIFIE une difference.

    La granularite « page » seule ne suffit pas, et c'est le piege qui a laisse
    passer l'incident : sur la page visee, seul le pied de page avait change, la
    synthese demandee etait intacte, et un controle par page aurait conclu a tort
    que la correction avait ete appliquee. Viser la page ET exclure le pied de
    page (exclure=r"^sources?\\s*:") redonne le bon verdict.
    """
    cibles = set(pages) if pages is not None else None
    re_motif = re.compile(motif, re.IGNORECASE) if motif else None
    re_excl = re.compile(exclure, re.IGNORECASE) if exclure else None

    def retenue(page, textes):
        if cibles is not None and page not in cibles:
            return False
        for t in textes:
            t = t.strip()
            if re_excl and re_excl.search(t):
                continue
            if re_motif and not re_motif.search(t):
                continue
            return True
        return False

    for m in res["modifies"]:
        for page in (m["page_apres"], m["page_avant"]):
            if retenue(page, [m["avant"], m["apres"]]):
                return True
    for x in res["ajoutes"] + res["retires"]:
        if retenue(x["page"], [x["texte"]]):
            return True
    return False


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    sortie = None
    if "--json" in sys.argv:
        sortie = sys.argv[sys.argv.index("--json") + 1]
        args = [a for a in args if a != sortie]
    if len(args) < 2:
        print("usage: block_diff.py <avant.pptx> <apres.pptx> [--json out.json]",
              file=sys.stderr)
        return 2
    res = comparer(args[0], args[1])
    if sortie:
        with open(sortie, "w", encoding="utf-8") as fh:
            json.dump(res, fh, ensure_ascii=False, indent=2)
    print("%d bloc(s) avant, %d apres, %d difference(s)"
          % (res["blocs_avant"], res["blocs_apres"], res["nb_differences"]))
    for l in rendu_humain(res):
        print("  - %s" % l)
    return 0 if res["nb_differences"] else 3


if __name__ == "__main__":
    sys.exit(main())
