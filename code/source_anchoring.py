#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Barriere d'ancrage des sources d'un document produit.

Defaut que ce controle ferme : un document genere cite en pied de page une
source qui n'a jamais ete deposee. Une barriere de sortie classique ne le voit
pas, parce qu'elle controle la forme (existence, taille, nombre de pages, traces
de fabrication, mise en page) et pas la veracite d'une attribution.

Principe, volontairement severe : une etiquette de source n'existe que si elle
est DECLAREE en configuration, et une etiquette adossee n'est autorisee que si la
piece qui la porte est reellement presente au depot. Toute autre etiquette est un
echec, avec le libelle fautif cite tel quel.

Contrat de forme attendu des pieds de page :

    Sources : <etiquette> · <etiquette> (precision) · <etiquette>

Chaque segment separe par « · » doit contenir une etiquette declaree. Un segment
sans etiquette declaree est refuse : c'est le seul moyen de refuser une
attribution inventee sans se fier au vocabulaire du redacteur.

Deux refus qui comptent autant que les autres :
  - aucune liste declaree en configuration -> echec (un controle qui se
    desactive quand sa configuration manque ne controle rien) ;
  - aucune mention de source dans le document -> echec (un rapport sans source
    citee n'est pas livrable).

Configuration attendue (YAML) :

    sources:
      adossees:
        - etiquette: "releve de caisses"
          pieces: ["releve-caisses", "M01"]     # motifs de nom de fichier
        - etiquette: "comptage de flux"
          pieces: ["comptage", "flux"]
      libres:
        - "commentaire de la direction"
        - "donnees publiques (verifiees par l'exploitant)"

Usage :
    source_anchoring.py <config.yaml> <repertoire_depot> <fichier.pptx> [...]
Sortie : lignes « OK|... » / « ECHEC|... », code retour 1 si un controle echoue.

Licence : MIT (voir LICENSE-CODE).
"""
import os
import re
import sys
import unicodedata

# Une ligne de sources peut apparaitre seule ou en fin de bloc de texte.
RE_ENTETE = re.compile(r"\b(sources?)\s*:\s*", re.IGNORECASE)
SEPARATEURS = re.compile(r"\s*[·•|]\s*")


def sansaccent(s):
    return "".join(c for c in unicodedata.normalize("NFD", s or "")
                   if unicodedata.category(c) != "Mn").lower().strip()


def normaliser(s):
    """Comparaison insensible aux accents, a la casse et aux espaces multiples."""
    return re.sub(r"\s+", " ", sansaccent(s))


# ------------------------------------------------------------- extraction
def blocs_du_deck(prs):
    """(numero_de_page, texte) pour chaque bloc de texte et chaque cellule."""
    out = []
    for i, slide in enumerate(prs.slides, start=1):
        for sh in slide.shapes:
            if sh.has_text_frame and sh.text_frame.text:
                out.append((i, sh.text_frame.text))
            if getattr(sh, "has_table", False) and sh.has_table:
                for row in sh.table.rows:
                    for cell in row.cells:
                        if cell.text:
                            out.append((i, cell.text))
    return out


def mentions_de_source(blocs):
    """-> [(page, segment_brut)] pour chaque etiquette citee."""
    trouvees = []
    for page, texte in blocs:
        for ligne in texte.splitlines():
            m = RE_ENTETE.search(ligne)
            if not m:
                continue
            queue = ligne[m.end():].strip()
            if not queue:
                continue
            for seg in SEPARATEURS.split(queue):
                seg = seg.strip(" .;")
                if seg:
                    trouvees.append((page, seg))
    return trouvees


# ---------------------------------------------------------------- ancrage
def charger_regles(cfg):
    """Liste fermee, telle que declaree en configuration.

    adossees : etiquette -> motifs de nom de piece qui l'adossent.
    libres   : etiquettes autorisees sans piece (emplacements reserves et
               sources publiques verifiees une par une par l'exploitant).
    """
    s = (cfg.get("sources") or {})
    adossees = []
    for e in s.get("adossees", []):
        adossees.append((normaliser(e["etiquette"]),
                         e["etiquette"],
                         [normaliser(p) for p in (e.get("pieces") or [])]))
    libres = [(normaliser(x), x) for x in s.get("libres", [])]
    return adossees, libres


def pieces_du_depot(dossier, ignores=()):
    """Noms des fichiers reellement deposes, normalises."""
    noms = []
    if dossier and os.path.isdir(dossier):
        for racine, _, fichiers in os.walk(dossier):
            for f in fichiers:
                if f.startswith(".") or f in ignores:
                    continue
                noms.append(normaliser(f))
    return noms


def controler(mentions, adossees, libres, pieces):
    """-> (liste d'echecs lisibles, liste de mentions validees)."""
    echecs, ok = [], []
    for page, brut in mentions:
        n = normaliser(brut)
        libre = next((orig for cle, orig in libres if cle and cle in n), None)
        if libre:
            ok.append((page, brut, "libre : %s" % libre))
            continue
        adossee = next(((cle, orig, motifs) for cle, orig, motifs in adossees
                        if cle and cle in n), None)
        if not adossee:
            echecs.append("page %d : source non declaree « %s »" % (page, brut))
            continue
        _, orig, motifs = adossee
        porteuse = next((p for p in pieces if any(m and m in p for m in motifs)), None)
        if porteuse:
            ok.append((page, brut, "adossee : %s" % orig))
        else:
            echecs.append("page %d : source « %s » citee sans piece correspondante au depot "
                          "(etiquette declaree « %s », aucune piece parmi %s)"
                          % (page, brut, orig, ", ".join(motifs) or "aucun motif"))
    return echecs, ok


def controler_fichier(chemin, cfg, dossier_depot):
    """-> (echecs, mentions_validees) pour un document."""
    from pptx import Presentation
    adossees, libres = charger_regles(cfg)
    if not adossees and not libres:
        return (["aucune liste de sources declaree en configuration : "
                 "l'ancrage des sources ne peut pas etre controle"], [])
    mentions = mentions_de_source(blocs_du_deck(Presentation(chemin)))
    if not mentions:
        return (["aucune mention de source trouvee dans le document : "
                 "un rapport sans source citee n'est pas livrable"], [])
    return controler(mentions, adossees, libres, pieces_du_depot(dossier_depot))


def main():
    import yaml
    if len(sys.argv) < 4:
        print(__doc__.strip().splitlines()[-4], file=sys.stderr)
        return 2
    cfg = yaml.safe_load(open(sys.argv[1], encoding="utf-8"))
    depot = sys.argv[2]
    rc = 0
    for chemin in sys.argv[3:]:
        base = os.path.basename(chemin)
        if not os.path.isfile(chemin):
            print("ECHEC|%s|fichier absent" % base)
            rc = 1
            continue
        echecs, ok = controler_fichier(chemin, cfg, depot)
        if echecs:
            rc = 1
            print("ECHEC|%s|source non adossee : %s" % (base, " · ".join(echecs)))
        else:
            print("OK|%s|%d mention(s) de source, toutes adossees au depot ou declarees"
                  % (base, len(ok)))
        for page, brut, motif in ok:
            print("    page %d : %s  [%s]" % (page, brut, motif))
    return rc


if __name__ == "__main__":
    sys.exit(main())
